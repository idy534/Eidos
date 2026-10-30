"""
Excel Formula Recalculation Script
Recalculates all formulas in an Excel file using LibreOffice
"""

import contextlib
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from office.convert import convert

from openpyxl import load_workbook

MAX_LOCATIONS = 100
EXTERNAL_REF_RE = re.compile(r"""(?<![\w"\[])'?\[\d+\][^!"\[\]]*'?!""")


def external_links_at_risk(filename):
    try:
        with zipfile.ZipFile(filename) as archive:
            names = archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return []
    if not any(n.startswith("xl/externalLinks/") for n in names):
        return []

    with contextlib.ExitStack() as stack:
        formulas = load_workbook(filename, data_only=False)
        stack.callback(formulas.close)
        values = load_workbook(filename, data_only=True)
        stack.callback(values.close)

        external_names = [
            name
            for name, dn in formulas.defined_names.items()
            if isinstance(getattr(dn, "value", None), str) and EXTERNAL_REF_RE.search(dn.value)
        ]
        name_re = (
            re.compile(r"\b(" + "|".join(re.escape(n) for n in external_names) + r")\b")
            if external_names
            else None
        )

        at_risk = []
        for sheet in formulas.sheetnames:
            ws = formulas[sheet]
            if not hasattr(ws, "iter_rows"):  
                continue
            cached = values[sheet]
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    reaches_out = EXTERNAL_REF_RE.search(v) or (name_re and name_re.search(v))
                    if reaches_out and cached[cell.coordinate].value is None:
                        at_risk.append(f"{sheet}!{cell.coordinate}")
        return at_risk


def recalc(filename, timeout=30, force=False):
    if not Path(filename).exists():
        return {"error": f"File {filename} does not exist"}

    abs_path = str(Path(filename).absolute())

    if not os.access(abs_path, os.W_OK):
        return {"error": f"{filename} is not writable; recalculation rewrites the file in place"}

    if not force:
        try:
            at_risk = external_links_at_risk(filename)
        except Exception as e:  
            return {"error": f"Could not inspect {filename} for external links: {e}"}
        if at_risk:
            shown = at_risk[:MAX_LOCATIONS]
            return {
                "error": (
                    "Refusing to recalculate: this workbook links to another workbook, and "
                    f"{len(at_risk)} linked cell(s) have lost their cached value (openpyxl strips "
                    "these on save). Recalculating would resolve them to #NAME? and delete the "
                    "external links for good. Copy those cells' values from the original file "
                    "before saving, or pass --force to accept the loss. Charts and conditional "
                    "formats can hold external references too, so this list may not be exhaustive."
                ),
                "external_link_cells": shown,
                "external_link_cells_truncated": max(0, len(at_risk) - len(shown)),
            }

    try:
        with tempfile.TemporaryDirectory(prefix="eidos-recalc-", dir=str(Path(abs_path).parent)) as directory:
            candidate = Path(directory) / Path(abs_path).name
            convert(abs_path, candidate, operation="recalculate", timeout=timeout)
            os.replace(candidate, abs_path)
    except (OSError, RuntimeError, subprocess.TimeoutExpired) as error:
        return {"error": f"Formula recalculation was not confirmed: {error}"}

    try:
        wb = load_workbook(filename, data_only=True)

        excel_errors = [
            "#VALUE!",
            "#DIV/0!",
            "#REF!",
            "#NAME?",
            "#NULL!",
            "#NUM!",
            "#N/A",
        ]
        error_details = {err: [] for err in excel_errors}
        total_errors = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if not hasattr(ws, "iter_rows"):  
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and isinstance(cell.value, str):
                        for err in excel_errors:
                            if err in cell.value:
                                location = f"{sheet_name}!{cell.coordinate}"
                                error_details[err].append(location)
                                total_errors += 1
                                break

        result = {
            "status": "success" if total_errors == 0 else "errors_found",
            "total_errors": total_errors,
            "error_summary": {},
        }

        for err_type, locations in error_details.items():
            if locations:
                entry = {"count": len(locations), "locations": locations[:MAX_LOCATIONS]}
                if len(locations) > MAX_LOCATIONS:
                    entry["locations_truncated"] = len(locations) - MAX_LOCATIONS
                result["error_summary"][err_type] = entry

        wb.close()

        wb_formulas = load_workbook(filename, data_only=False)
        formula_count = 0
        for sheet_name in wb_formulas.sheetnames:
            ws = wb_formulas[sheet_name]
            if not hasattr(ws, "iter_rows"):  
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if (
                        cell.value
                        and isinstance(cell.value, str)
                        and cell.value.startswith("=")
                    ):
                        formula_count += 1
        wb_formulas.close()

        result["total_formulas"] = formula_count

        return result

    except Exception as e:
        return {"error": str(e)}


def main():
    args = [a for a in sys.argv[1:] if a != "--force"]
    force = "--force" in sys.argv[1:]

    if not args:
        print("Usage: python recalc.py <excel_file> [timeout_seconds] [--force]")
        print("\nRecalculates all formulas in an Excel file using LibreOffice")
        print("\nReturns JSON with error details:")
        print("  - status: 'success' or 'errors_found'")
        print("  - total_errors: Total number of Excel errors found")
        print("  - total_formulas: Number of formulas in the file")
        print("  - error_summary: Breakdown by error type with locations")
        print("    - #VALUE!, #DIV/0!, #REF!, #NAME?, #NULL!, #NUM!, #N/A")
        print("\nOn any failure the JSON has an 'error' key and no 'status'.")
        print("--force recalculates even when it would destroy external links.")
        sys.exit(1)

    filename = args[0]
    timeout = int(args[1]) if len(args) > 1 else 30

    result = recalc(filename, timeout, force=force)
    print(json.dumps(result, indent=2))
    sys.exit(1 if "error" in result else 0)


if __name__ == "__main__":
    main()
