"""Print a bounded range of populated cells, preserving formula text."""
import argparse

from openpyxl import load_workbook

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("file")
parser.add_argument("--start-row", type=int, default=1)
parser.add_argument("--rows", type=int, default=100)
args = parser.parse_args()
if args.start_row < 1 or not 1 <= args.rows <= 1000:
    parser.error("start-row must be positive; rows must be between 1 and 1000")
workbook = load_workbook(args.file, read_only=True, data_only=False)
try:
    for sheet in workbook:
        print(f"## {sheet.title}")
        for row in sheet.iter_rows(min_row=args.start_row, max_row=args.start_row + args.rows - 1):
            print(" | ".join(f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None))
finally:
    workbook.close()
